"""
DPR financials → Excel workbook exporter.

Per KAU pre-UAT reply §6.3 (2026-09-08): bankers and appraisal officers need
the projected financial tables in a spreadsheet, not just a PDF, so they can
verify the calc engine's numbers cell-by-cell against their own models.

Public entry point:
    * `render_financials_workbook(project) -> BytesIO`

Produced sheets (fixed order — banker-familiar):
    A. Project Summary       — meta + top-line ratios
    B. Cost & MoF Variance   — user vs derived + gap
    C. Capital Schedule      — per-month capex plan
    D. Depreciation          — per asset class + per year
    E. Interest Schedule     — per year amortisation
    F. Profit & Loss         — Y1..YN
    G. Cash Flow             — Y0..YN
    H. Balance Sheet         — Y0..YN + invariant check
    I. Ratios                — NPV / IRR / DSCR / payback / break-even
    J. DSCR per year         — year-by-year debt service coverage

Author: Athul Gopan (Kefi Tech Solutions)
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Iterable, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.fpo.services.dpr.calculation import CalculationResult, compute


# ── Styling constants ─────────────────────────────────────────────────────
_HEADER_FILL = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
_HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
_SECTION_FILL = PatternFill(start_color='E86C1A', end_color='E86C1A', fill_type='solid')
_SECTION_FONT = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
_ROW_ALT_FILL = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
_NUM_FMT = '#,##0.00'
_INT_FMT = '#,##0'
_PCT_FMT = '0.00"%"'


def render_financials_workbook(project) -> BytesIO:
    """Compute the DPR then materialise all 10 sheets. Returns an in-memory
    BytesIO with the .xlsx bytes ready to stream via `FileResponse`.
    """
    result: CalculationResult = compute(project)
    wb = openpyxl.Workbook()

    # openpyxl creates a default sheet — repurpose it for the summary
    # instead of leaving an empty first tab.
    _write_summary_sheet(wb.active, project, result)

    _write_cost_mof_sheet(wb.create_sheet('B — Cost & MoF'), result)
    _write_capital_schedule_sheet(wb.create_sheet('C — Capital Schedule'), result)
    _write_depreciation_sheet(wb.create_sheet('D — Depreciation'), result)
    _write_interest_sheet(wb.create_sheet('E — Interest'), result)
    _write_pnl_sheet(wb.create_sheet('F — P&L'), result)
    _write_cash_flow_sheet(wb.create_sheet('G — Cash Flow'), result)
    _write_balance_sheet_sheet(wb.create_sheet('H — Balance Sheet'), result)
    _write_ratios_sheet(wb.create_sheet('I — Ratios'), result)
    _write_dscr_sheet(wb.create_sheet('J — DSCR by Year'), result)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# ─────────────────────────────────────────────────────────────────────────────
# Sheet writers
# ─────────────────────────────────────────────────────────────────────────────

def _write_summary_sheet(ws, project, r: CalculationResult) -> None:
    ws.title = 'A — Summary'
    _set_col_widths(ws, [34, 24])

    _write_section_title(ws, 1, 1, 'DPR Financial Summary')
    ws.cell(row=2, column=1, value='FPO name').font = Font(bold=True)
    ws.cell(row=2, column=2, value=getattr(getattr(project, 'fpo', None), 'name', '') or '—')

    ws.cell(row=3, column=1, value='Project title').font = Font(bold=True)
    ws.cell(row=3, column=2, value=project.title or '—')

    ws.cell(row=4, column=1, value='Project UUID').font = Font(bold=True)
    ws.cell(row=4, column=2, value=str(project.uuid))

    ws.cell(row=5, column=1, value='Projection horizon').font = Font(bold=True)
    ws.cell(row=5, column=2, value=f'{r.projection_years} years')

    ws.cell(row=6, column=1, value='Total project cost').font = Font(bold=True)
    _num_cell(ws, 6, 2, r.cost.total)

    ws.cell(row=7, column=1, value='Total means of finance').font = Font(bold=True)
    _num_cell(ws, 7, 2, r.mof.total)

    ws.cell(row=8, column=1, value='MoF − Cost delta').font = Font(bold=True)
    _num_cell(ws, 8, 2, r.variance.delta)
    ws.cell(row=9, column=1, value='Variance %').font = Font(bold=True)
    _pct_cell(ws, 9, 2, r.variance.pct)

    # Ratios strip — same order the PDF cover surfaces them.
    _write_section_title(ws, 11, 1, 'Key Ratios')
    ratios = r.ratios
    row = 12
    if ratios:
        ws.cell(row=row, column=1, value='Discount rate').font = Font(bold=True)
        _pct_cell(ws, row, 2, ratios.discount_rate_pct); row += 1
        ws.cell(row=row, column=1, value='NPV (₹)').font = Font(bold=True)
        _num_cell(ws, row, 2, ratios.npv); row += 1
        ws.cell(row=row, column=1, value='IRR (%)').font = Font(bold=True)
        _pct_cell(ws, row, 2, ratios.irr_pct); row += 1
        ws.cell(row=row, column=1, value='Min DSCR').font = Font(bold=True)
        _num_cell(ws, row, 2, ratios.dscr_min); row += 1
        ws.cell(row=row, column=1, value='Avg DSCR').font = Font(bold=True)
        _num_cell(ws, row, 2, ratios.dscr_avg); row += 1
        ws.cell(row=row, column=1, value='Payback (years)').font = Font(bold=True)
        _num_cell(ws, row, 2, ratios.payback_period_years); row += 1
        ws.cell(row=row, column=1, value='Break-even year').font = Font(bold=True)
        ws.cell(row=row, column=2, value=ratios.break_even_year if ratios.break_even_year is not None else '—')
    else:
        ws.cell(row=row, column=1, value='(Ratios not available — insufficient data)')


def _write_cost_mof_sheet(ws, r: CalculationResult) -> None:
    _set_col_widths(ws, [40, 22, 22, 22])
    _write_section_title(ws, 1, 1, 'A. Project Cost — Line-item Breakdown')
    _write_header(ws, 2, ['Cost line', 'Amount (₹)'])
    row = 3
    for k, v in sorted(r.cost.by_field.items()):
        ws.cell(row=row, column=1, value=k)
        _num_cell(ws, row, 2, v); row += 1
    ws.cell(row=row, column=1, value='TOTAL').font = Font(bold=True)
    _num_cell(ws, row, 2, r.cost.total); row += 2

    _write_section_title(ws, row, 1, 'B. Means of Finance — Line-item Breakdown')
    row += 1
    _write_header(ws, row, ['MoF line', 'Amount (₹)']); row += 1
    for k, v in sorted(r.mof.by_field.items()):
        ws.cell(row=row, column=1, value=k)
        _num_cell(ws, row, 2, v); row += 1
    ws.cell(row=row, column=1, value='TOTAL').font = Font(bold=True)
    _num_cell(ws, row, 2, r.mof.total); row += 2

    _write_section_title(ws, row, 1, 'C. Cost ↔ MoF Reconciliation')
    row += 1
    ws.cell(row=row, column=1, value='Total project cost').font = Font(bold=True)
    _num_cell(ws, row, 2, r.variance.cost_total); row += 1
    ws.cell(row=row, column=1, value='Total MoF').font = Font(bold=True)
    _num_cell(ws, row, 2, r.variance.mof_total); row += 1
    ws.cell(row=row, column=1, value='MoF − Cost delta').font = Font(bold=True)
    _num_cell(ws, row, 2, r.variance.delta); row += 1
    ws.cell(row=row, column=1, value='Variance %').font = Font(bold=True)
    _pct_cell(ws, row, 2, r.variance.pct); row += 1
    ws.cell(row=row, column=1, value='Threshold %').font = Font(bold=True)
    _pct_cell(ws, row, 2, r.variance.threshold_pct); row += 1
    ws.cell(row=row, column=1, value='Exceeds threshold?').font = Font(bold=True)
    ws.cell(row=row, column=2, value='Yes' if r.variance.exceeds_threshold else 'No')


def _write_capital_schedule_sheet(ws, r: CalculationResult) -> None:
    _set_col_widths(ws, [10, 18, 18, 20, 20, 20])
    sched = r.capital_schedule
    if not sched:
        ws.cell(row=1, column=1, value='(Capital schedule not available — no tranches recorded)')
        return
    _write_section_title(ws, 1, 1, 'Capital Schedule (per-month capex plan)')
    ws.cell(row=2, column=1, value=sched.distribution_note).font = Font(italic=True)
    ws.cell(row=3, column=1, value=f'is_estimated: {"Yes" if sched.is_estimated else "No"} · '
                                    f'implementation period: {sched.implementation_period_months} months').font = Font(italic=True)
    _write_header(ws, 5, ['Month', 'Cost incurred', 'MoF received',
                          'Cum cost', 'Cum MoF', 'Unfunded'])
    row = 6
    for r_row in sched.rows:
        ws.cell(row=row, column=1, value=f'M{r_row.month}')
        _num_cell(ws, row, 2, r_row.cost_incurred)
        _num_cell(ws, row, 3, r_row.mof_received)
        _num_cell(ws, row, 4, r_row.cumulative_cost)
        _num_cell(ws, row, 5, r_row.cumulative_mof)
        _num_cell(ws, row, 6, r_row.unfunded_balance)
        row += 1
    ws.cell(row=row, column=1, value='FINAL').font = Font(bold=True)
    _num_cell(ws, row, 4, sched.final_cost)
    _num_cell(ws, row, 5, sched.final_mof)


def _write_depreciation_sheet(ws, r: CalculationResult) -> None:
    dep = r.depreciation
    if not dep:
        ws.cell(row=1, column=1, value='(Depreciation schedule not available)')
        return
    _write_section_title(ws, 1, 1, 'Depreciation Schedule — SLM, per asset class')
    years = list(range(1, r.projection_years + 1))
    header = ['Asset class', 'Initial cost (₹)', 'Rate (%)', *[f'Y{y} dep' for y in years], 'Total dep']
    _set_col_widths(ws, [22, 20, 12] + [14] * len(years) + [20])
    _write_header(ws, 2, header)
    row = 3
    for cls in dep.classes:
        ws.cell(row=row, column=1, value=cls.label or cls.key)
        _num_cell(ws, row, 2, cls.initial_cost)
        _pct_cell(ws, row, 3, cls.rate_pct)
        col = 4
        cls_rows = {ar.year: ar.depreciation for ar in cls.rows}
        for y in years:
            _num_cell(ws, row, col, cls_rows.get(y, Decimal('0')))
            col += 1
        _num_cell(ws, row, col, sum((ar.depreciation for ar in cls.rows), Decimal('0')))
        row += 1
    # Yearly totals (already computed on dep.total_depreciation_by_year)
    ws.cell(row=row, column=1, value='TOTAL').font = Font(bold=True)
    col = 4
    for y in years:
        _num_cell(ws, row, col, dep.total_depreciation_by_year.get(y, Decimal('0')))
        col += 1


def _write_interest_sheet(ws, r: CalculationResult) -> None:
    isch = r.interest_schedule
    _set_col_widths(ws, [8, 22, 22, 22, 22])
    _write_section_title(ws, 1, 1, 'Interest & Principal Amortisation Schedule')
    if not isch or not isch.rows:
        ws.cell(row=2, column=1, value='(No loan proposed — schedule empty)')
        return
    ws.cell(row=2, column=1, value=f'Method: {isch.repayment_method}').font = Font(italic=True)
    ws.cell(row=3, column=1, value=f'Moratorium: {isch.moratorium_months} months · '
                                    f'{isch.moratorium_interest_treatment}').font = Font(italic=True)
    _write_header(ws, 5, ['Year', 'Opening bal', 'Interest', 'Principal', 'Closing bal'])
    row = 6
    for r_row in isch.rows:
        ws.cell(row=row, column=1, value=r_row.year)
        _num_cell(ws, row, 2, r_row.opening_balance)
        _num_cell(ws, row, 3, r_row.interest)
        _num_cell(ws, row, 4, r_row.principal)
        _num_cell(ws, row, 5, r_row.closing_balance)
        row += 1


def _write_pnl_sheet(ws, r: CalculationResult) -> None:
    pl = r.profit_loss
    _set_col_widths(ws, [22, 8] + [16] * 10)
    _write_section_title(ws, 1, 1, 'Projected Profit & Loss (₹)')
    if not pl:
        ws.cell(row=2, column=1, value='(P&L not available)')
        return
    years = [row.year for row in pl.rows]
    _write_header(ws, 2, ['Line item', ''] + [f'Y{y}' for y in years])

    lines = [
        ('Revenue', 'revenue'),
        ('Operating cost', 'operating_cost'),
        ('EBITDA', 'ebitda'),
        ('Depreciation', 'depreciation'),
        ('EBIT', 'ebit'),
        ('Interest', 'interest'),
        ('PBT', 'pbt'),
        ('Tax', 'tax'),
        ('PAT', 'pat'),
    ]
    row = 3
    for label, attr in lines:
        ws.cell(row=row, column=1, value=label).font = Font(bold=(label in ('EBITDA', 'EBIT', 'PBT', 'PAT')))
        col = 3
        for r_row in pl.rows:
            _num_cell(ws, row, col, getattr(r_row, attr))
            col += 1
        row += 1

    # Cumulative PAT row for banker sight-line
    row += 1
    ws.cell(row=row, column=1, value='Cumulative PAT').font = Font(bold=True, italic=True)
    col = 3
    for y in years:
        _num_cell(ws, row, col, pl.cumulative_pat_by_year.get(y, Decimal('0')))
        col += 1


def _write_cash_flow_sheet(ws, r: CalculationResult) -> None:
    cf = r.cash_flow
    _set_col_widths(ws, [30, 8] + [16] * 11)
    _write_section_title(ws, 1, 1, 'Projected Cash Flow (₹)')
    if not cf:
        ws.cell(row=2, column=1, value='(Cash flow not available)')
        return
    years = [row.year for row in cf.rows]
    _write_header(ws, 2, ['Line item', ''] + [f'Y{y}' for y in years])

    lines = [
        ('PAT', 'pat'),
        ('Depreciation add-back', 'depreciation_addback'),
        ('WC change', 'working_capital_change'),
        ('Cash from operations', 'cash_from_operations'),
        ('Capex', 'capex'),
        ('Cash from investing', 'cash_from_investing'),
        ('MoF inflow', 'mof_inflow'),
        ('Loan principal repayment', 'loan_principal_repayment'),
        ('Cash from financing', 'cash_from_financing'),
        ('Net cash flow', 'net_cash_flow'),
        ('Opening cash', 'opening_cash'),
        ('Closing cash', 'closing_cash'),
    ]
    row = 3
    for label, attr in lines:
        bold = label in (
            'Cash from operations', 'Cash from investing', 'Cash from financing',
            'Net cash flow', 'Closing cash',
        )
        ws.cell(row=row, column=1, value=label).font = Font(bold=bold)
        col = 3
        for r_row in cf.rows:
            _num_cell(ws, row, col, getattr(r_row, attr))
            col += 1
        row += 1

    row += 1
    ws.cell(row=row, column=1, value='Ending cash').font = Font(bold=True, italic=True)
    _num_cell(ws, row, 3, cf.ending_cash)


def _write_balance_sheet_sheet(ws, r: CalculationResult) -> None:
    bs = r.balance_sheet
    _set_col_widths(ws, [32, 8] + [16] * 11)
    _write_section_title(ws, 1, 1, 'Projected Balance Sheet (₹)')
    if not bs:
        ws.cell(row=2, column=1, value='(Balance sheet not available)')
        return
    years = [row.year for row in bs.rows]
    _write_header(ws, 2, ['Line item', ''] + [f'Y{y}' for y in years])

    lines = [
        ('— ASSETS —', None),
        ('Land', 'land'),
        ('CWIP', 'cwip'),
        ('Gross fixed assets', 'gross_fixed_assets'),
        ('Accumulated depreciation', 'accumulated_depreciation'),
        ('Net fixed assets', 'net_fixed_assets'),
        ('Working capital', 'working_capital'),
        ('Cash & bank', 'cash_and_bank'),
        ('TOTAL ASSETS', 'total_assets'),
        ('— EQUITY & LIABILITIES —', None),
        ('Promoter equity', 'promoter_equity'),
        ('Capital reserve', 'capital_reserve'),
        ('Retained earnings', 'retained_earnings'),
        ('Total equity', 'total_equity'),
        ('Term loan outstanding', 'term_loan_outstanding'),
        ('Other liabilities', 'other_liabilities'),
        ('Total liabilities', 'total_liabilities'),
        ('TOTAL EQUITY & LIABILITIES', 'total_equity_and_liabilities'),
        ('Invariant delta (should ≈ 0)', 'invariant_delta'),
        ('Invariant OK', 'invariant_ok'),
    ]
    row = 3
    for label, attr in lines:
        cell = ws.cell(row=row, column=1, value=label)
        if attr is None:
            cell.font = Font(bold=True)
            cell.fill = _ROW_ALT_FILL
            row += 1
            continue
        cell.font = Font(bold=label.startswith('TOTAL') or label.startswith('Total'))
        col = 3
        for r_row in bs.rows:
            v = getattr(r_row, attr)
            if isinstance(v, bool):
                ws.cell(row=row, column=col, value='Yes' if v else 'No')
            else:
                _num_cell(ws, row, col, v)
            col += 1
        row += 1
    row += 1
    ws.cell(row=row, column=1, value='All years balanced?').font = Font(bold=True)
    ws.cell(row=row, column=3, value='Yes' if bs.all_years_balanced else 'No')
    ws.cell(row=row + 1, column=1, value='Max invariant delta').font = Font(bold=True)
    _num_cell(ws, row + 1, 3, bs.max_invariant_delta)


def _write_ratios_sheet(ws, r: CalculationResult) -> None:
    _set_col_widths(ws, [32, 22])
    _write_section_title(ws, 1, 1, 'Bankability Ratios')
    ratios = r.ratios
    if not ratios:
        ws.cell(row=2, column=1, value='(Ratios not available)')
        return
    row = 2
    def _kv(label, value, fmt='num'):
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        if value is None:
            ws.cell(row=row, column=2, value='—')
        elif fmt == 'pct':
            _pct_cell(ws, row, 2, value)
        elif fmt == 'int':
            ws.cell(row=row, column=2, value=int(value)).number_format = _INT_FMT
        elif fmt == 'text':
            ws.cell(row=row, column=2, value=str(value))
        else:
            _num_cell(ws, row, 2, value)
        row += 1

    _kv('Discount rate (%)', ratios.discount_rate_pct, 'pct')
    _kv('NPV (₹)', ratios.npv)
    _kv('IRR (%)', ratios.irr_pct, 'pct')
    _kv('IRR bisection converged', 'Yes' if ratios.irr_converged else 'No', 'text')
    _kv('Min DSCR', ratios.dscr_min)
    _kv('Avg DSCR', ratios.dscr_avg)
    _kv('Payback (years)', ratios.payback_period_years)
    _kv('Break-even year', ratios.break_even_year, 'int')


def _write_dscr_sheet(ws, r: CalculationResult) -> None:
    _set_col_widths(ws, [8, 22, 22, 16])
    _write_section_title(ws, 1, 1, 'Debt Service Coverage Ratio — Year by Year')
    ratios = r.ratios
    if not ratios or not ratios.dscr_rows:
        ws.cell(row=2, column=1, value='(No DSCR rows — likely no debt service)')
        return
    _write_header(ws, 2, ['Year', 'Numerator (PAT+Dep+Int)', 'Denominator (Int+Prin)', 'DSCR'])
    row = 3
    for r_row in ratios.dscr_rows:
        ws.cell(row=row, column=1, value=r_row.year)
        _num_cell(ws, row, 2, r_row.numerator)
        _num_cell(ws, row, 3, r_row.denominator)
        if r_row.dscr is None:
            ws.cell(row=row, column=4, value='—')
        else:
            _num_cell(ws, row, 4, r_row.dscr)
        row += 1


# ─────────────────────────────────────────────────────────────────────────────
# Cell helpers
# ─────────────────────────────────────────────────────────────────────────────

def _set_col_widths(ws, widths: Iterable[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_section_title(ws, row: int, col: int, title: str) -> None:
    cell = ws.cell(row=row, column=col, value=title)
    cell.font = _SECTION_FONT
    cell.fill = _SECTION_FILL
    cell.alignment = Alignment(vertical='center')


def _write_header(ws, row: int, headers: list[str]) -> None:
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _num_cell(ws, row: int, col: int, value: Optional[Decimal]) -> None:
    if value is None:
        ws.cell(row=row, column=col, value='')
        return
    n = float(value) if isinstance(value, Decimal) else value
    cell = ws.cell(row=row, column=col, value=n)
    cell.number_format = _NUM_FMT
    cell.alignment = Alignment(horizontal='right')


def _pct_cell(ws, row: int, col: int, value: Optional[Decimal]) -> None:
    if value is None:
        ws.cell(row=row, column=col, value='—')
        return
    n = float(value) if isinstance(value, Decimal) else value
    cell = ws.cell(row=row, column=col, value=n)
    cell.number_format = _PCT_FMT
    cell.alignment = Alignment(horizontal='right')
