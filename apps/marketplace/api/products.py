"""
Arunima S

Product API — FPO Product Listings
===================================
Base Path: /api/marketplace/products/
"""
import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone

from drf_spectacular.utils import extend_schema, extend_schema_view
from rest_framework.decorators import action
from rest_framework.parsers import MultiPartParser

from apps.core.exceptions import BusinessLogicError
from apps.core.models.generic import MasterLookup
from apps.core.permissions.rbac import IsAuthenticated, IsFPOManager
from apps.core.services.translation import t
from apps.core.utils.pagination import StandardPagination
from apps.core.utils.responses import StandardResponse
from apps.core.views import TranslatedViewSet
from apps.core.services.fpo_permission import get_member_fpo
from apps.database.models import Product, ProductStock
from apps.marketplace.permissions import CanManageProducts, IsApprovedFPO
from apps.marketplace.serializers import ProductSerializer
from apps.marketplace.services import (
    activate_would_leave_all_inactive, run_matching, set_product_active, switch_product_on_for_live_batch,
)


def _clear_public_market_cache():
    """
    Wipe all cached Public Market Hub product-list responses.
    Called whenever a product's status/public visibility could change,
    so new/updated products show up immediately instead of waiting for
    the 1-hour cache TTL to expire naturally.
    """
    from django.core.cache import cache
    try:
        cache.delete_pattern('public:market:products:*')
        cache.delete_pattern('public:market:commodities:*')
        cache.delete_pattern('public:market:opportunities')
    except AttributeError:
        # Cache backend doesn't support delete_pattern (e.g. LocMemCache in tests)
        pass


# =============================================================================
# BULK IMPORT — parsing/validation helpers for POST .../bulk-import/
# =============================================================================

def _parse_decimal(value, field_name):
    """Parse a required positive-number cell (quantity / price_per_unit)."""
    if value is None or str(value).strip() == '':
        raise ValueError(f'{field_name} is required')
    try:
        parsed = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"{field_name} must be a number, got '{value}'")
    if parsed <= 0:
        raise ValueError(f'{field_name} must be greater than 0')
    return parsed


def _parse_bool(value, default: bool) -> bool:
    """
    Parse a loosely-formatted boolean cell. Accepts common Yes/No, True/False,
    1/0, Y/N variants (case-insensitive) and treats blank as `default` so
    older template downloads (which omit this column) keep working.
    """
    if value is None:
        return default
    text = str(value).strip().lower()
    if text == "":
        return default
    if text in {"true", "yes", "y", "1"}:
        return True
    if text in {"false", "no", "n", "0"}:
        return False
    raise ValueError(f"must be Yes/No (got '{value}')")


def _parse_date(value, field_name, required=False):
    """
    Parse a date cell. openpyxl hands back a real datetime/date object for
    date-formatted Excel cells; a CSV upload hands back a plain string
    ('YYYY-MM-DD'), so both are accepted here.
    """
    if value is None or str(value).strip() == '':
        if required:
            raise ValueError(f'{field_name} is required')
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), '%Y-%m-%d').date()
    except ValueError:
        raise ValueError(f"{field_name} must be a valid date (YYYY-MM-DD), got '{value}'")


def _normalize_product_name(name):
    """
    Grouping key for a product name — trimmed, inner whitespace collapsed and
    case-insensitive, so 'Organic Rice' and ' organic  rice' are one product.
    """
    return ' '.join(str(name or '').split()).casefold()


# The batch values a sheet row carries. Two batches of the same product with
# equal values here are duplicates (the bulk import skips the second one).
_BATCH_SIGNATURE_FIELDS = (
    'quantity', 'unit', 'price_per_unit', 'quality_certification',
    'available_from', 'available_until', 'is_public',
)


def _batch_signature(values):
    """Hashable signature of a batch, from a parsed row's batch dict or a ProductStock."""
    if isinstance(values, dict):
        return tuple(values[f] for f in _BATCH_SIGNATURE_FIELDS)
    return tuple(getattr(values, f) for f in _BATCH_SIGNATURE_FIELDS)


def _parse_bulk_product(row, commodity_cache=None):
    """
    Validate the product columns of one bulk-import row (name, commodity,
    description) and return them cleaned. Nothing is written here.

    Raises ValueError with a human-readable message on any validation
    failure; the caller catches this and skips just this row, so one bad
    row never blocks the rest of the file.
    """
    name_en = str(row.get('name_en') or '').strip()
    if not name_en:
        raise ValueError('name_en is required')
    name_ml = str(row.get('name_ml') or '').strip()

    commodity_code = str(row.get('commodity_code') or '').strip().lower()
    if not commodity_code:
        raise ValueError('commodity_code is required')
    if commodity_cache is not None and commodity_code in commodity_cache:
        commodity = commodity_cache[commodity_code]
    else:
        try:
            commodity = MasterLookup.objects.get(
                category='commodity', code=commodity_code, is_active=True,
            )
        except MasterLookup.DoesNotExist:
            raise ValueError(f"Unknown commodity_code '{commodity_code}' — check the Commodity Codes sheet")
        if commodity_cache is not None:
            commodity_cache[commodity_code] = commodity

    description_en = str(row.get('description_en') or '').strip()
    if not description_en:
        raise ValueError('description_en is required')
    description_ml = str(row.get('description_ml') or '').strip()

    return {
        'name_en': name_en,
        'name_ml': name_ml,
        'commodity': commodity,
        'description_en': description_en,
        'description_ml': description_ml,
    }


def _parse_bulk_batch(row):
    """
    Validate the batch columns of one bulk-import row and return them as
    ProductStock fields. Kept apart from _parse_bulk_product so a row with
    good product details but a bad batch can still create the product.

    Raises ValueError with a human-readable message on the first problem.
    """
    quantity = _parse_decimal(row.get('quantity'), 'quantity')
    price_per_unit = _parse_decimal(row.get('price_per_unit'), 'price_per_unit')

    unit = str(row.get('unit') or '').strip().lower()
    valid_units = [choice[0] for choice in ProductStock.Unit.choices]
    if unit not in valid_units:
        raise ValueError(f"unit must be one of {', '.join(valid_units)}, got '{unit}'")

    quality_certification = str(row.get('quality_certification') or '').strip()

    available_from = _parse_date(row.get('available_from'), 'available_from', required=True)
    available_until = _parse_date(row.get('available_until'), 'available_until', required=False)
    # Same rules as the Add Batch form — imported batches go live immediately.
    if available_until and available_until < available_from:
        raise ValueError('available_until must be the same as or after available_from')
    if available_until and available_until < timezone.localdate():
        raise ValueError('available_until cannot be in the past')

    # Optional is_public column — default True so bulk-imported batches are
    # discoverable on the public Market Hub right away (matches the default
    # the Add Product form picks). The FPO can set FALSE / NO / 0 in the
    # cell to opt a specific row out.
    try:
        is_public = _parse_bool(row.get('is_public'), default=True)
    except ValueError as e:
        raise ValueError(f'is_public {e}')

    return {
        'quantity': quantity,
        'unit': unit,
        'price_per_unit': price_per_unit,
        'quality_certification': quality_certification,
        'available_from': available_from,
        'available_until': available_until,
        'is_public': is_public,
    }


@extend_schema_view(
    list=extend_schema(tags=['Marketplace - Products']),
    create=extend_schema(tags=['Marketplace - Products']),
    retrieve=extend_schema(tags=['Marketplace - Products']),
    update=extend_schema(tags=['Marketplace - Products']),
    partial_update=extend_schema(tags=['Marketplace - Products']),
    destroy=extend_schema(tags=['Marketplace - Products']),
)
class ProductViewSet(TranslatedViewSet):
    """
    GET/POST          /api/marketplace/products/
    GET/PATCH/DELETE  /api/marketplace/products/{id}/
    POST              /api/marketplace/products/{id}/publish/
    POST              /api/marketplace/products/{id}/mark-sold/
    GET               /api/marketplace/products/bulk-template/
    POST              /api/marketplace/products/bulk-import/
    """

    serializer_class = ProductSerializer
    permission_classes = [IsAuthenticated, IsFPOManager, IsApprovedFPO, CanManageProducts]
    pagination_class = StandardPagination

    list_message = 'marketplace.products_retrieved'
    create_message = 'marketplace.product_created'
    update_message = 'marketplace.product_updated'
    destroy_message = 'marketplace.product_deleted'

    def get_queryset(self):
        from django.db.models import Q

        # FPO members only see their own FPO's products; exclude soft-deleted rows.
        # get_member_fpo() resolves to the FPO for both primary + secondary
        # users (via FPOUserMembership) so team members can also browse.
        fpo = get_member_fpo(self.request.user)
        queryset = (
            Product.objects.filter(fpo=fpo, is_deleted=False)
            .select_related('commodity', 'fpo')
            .prefetch_related('stocks')
            .order_by('-created_at')
        )

        search = self.request.query_params.get('search', '').strip()
        if search:
            # Batch numbers are searchable too (e.g. "P34-B2" from an inquiry).
            queryset = queryset.filter(
                Q(name__en__icontains=search) | Q(name__ml__icontains=search)
                | Q(stocks__batch_number__icontains=search)
            ).distinct()

        status = self.request.query_params.get('status')
        if status in (
            ProductStock.Status.DRAFT, ProductStock.Status.ACTIVE, ProductStock.Status.INACTIVE,
            ProductStock.Status.SOLD, ProductStock.Status.EXPIRED,
        ):
            # Any batch in the requested status qualifies the product. distinct()
            # because the inner JOIN duplicates products with multiple matching
            # batches.
            queryset = queryset.filter(stocks__status=status).distinct()

        # Product-level switch (the products table's Status column / filter).
        product_status = self.request.query_params.get('product_status')
        if product_status in ('active', 'inactive'):
            queryset = queryset.filter(is_active=product_status == 'active')
        return queryset

    # update()/destroy() are NOT overridden — per house convention, business
    # rule checks live in perform_update()/perform_destroy(), raising
    # BusinessLogicError (HTTP 422). custom_exception_handler (apps/core/
    # exceptions/handlers.py) catches BaseAPIException subclasses and formats
    # them into the standard {"status": "error", "message": ..., "code": ...}
    # shape via exc.to_dict() — so TranslatedViewSet's built-in update()/
    # destroy() (which already return StandardResponse on success) stay
    # untouched.
    #
    # NOTE: BusinessLogicError returns HTTP 422, not 400 — this is a genuine
    # behavior change from the earlier StandardResponse.error(..., status_code=400)
    # version. 422 is arguably more correct for "valid request, wrong state"
    # (vs 400 "malformed request"), but flag this if the frontend specifically
    # expects 400 for these cases.
    def perform_create(self, serializer):
        serializer.save()
        _clear_public_market_cache()

    def perform_update(self, serializer):
        # Product-master edits (name / commodity / description / image) are
        # always safe — stock lifecycle restrictions live on the nested
        # /stocks/ endpoint instead.
        serializer.save()
        _clear_public_market_cache()

    def perform_destroy(self, instance):
        # Soft delete — only if no non-draft stock batches exist. BaseModel
        # provides .soft_delete(), which sets is_deleted=True instead of
        # removing the row. A product with ACTIVE/SOLD/EXPIRED batches
        # can't be deleted — those must be ended through the stock endpoint
        # first.
        non_draft = instance.stocks.exclude(status=ProductStock.Status.DRAFT).exists()
        if non_draft:
            raise BusinessLogicError(
                message=t('marketplace.only_draft_deletable', self.get_language()),
                code='only_draft_deletable',
            )
        instance.soft_delete(user=self.request.user)
        _clear_public_market_cache()

    @extend_schema(tags=['Marketplace - Products'])
    @action(detail=True, methods=['post'])
    def publish(self, request, pk=None):
        """
        Legacy convenience action — publishes the product's single DRAFT
        batch. Only works when the product has exactly one DRAFT stock
        (the common case right after Product+Stock were created together
        via POST /products/). For multi-batch workflows, use the nested
        POST /products/{id}/stocks/{stock_id}/publish/ action instead.
        """
        product = self.get_object()
        drafts = list(product.stocks.filter(status=ProductStock.Status.DRAFT))
        if len(drafts) != 1:
            raise BusinessLogicError(
                message=t('marketplace.only_draft_publishable', self.get_language()),
                code='only_draft_publishable',
            )
        stock = drafts[0]
        if stock.available_until and stock.available_until < timezone.localdate():
            raise BusinessLogicError(
                message=t('marketplace.batch_validity_ended', self.get_language()),
                code='batch_validity_ended',
            )
        stock.status = ProductStock.Status.ACTIVE
        stock.save()
        switch_product_on_for_live_batch(product)

        run_matching(stock)
        _clear_public_market_cache()

        return StandardResponse.success(
            data=ProductSerializer(product).data,
            message=t('marketplace.product_published', self.get_language())
        )

    @extend_schema(tags=['Marketplace - Products'])
    @action(detail=True, methods=['post'], url_path='mark-sold')
    def mark_sold(self, request, pk=None):
        """
        Legacy convenience action — marks the product's single ACTIVE
        batch as SOLD. Only works when exactly one ACTIVE batch exists.
        For multi-batch workflows, use the nested
        POST /products/{id}/stocks/{stock_id}/mark-sold/ action instead.
        """
        product = self.get_object()
        actives = list(product.stocks.filter(status=ProductStock.Status.ACTIVE))
        if len(actives) != 1:
            raise BusinessLogicError(
                message=t('marketplace.only_active_can_be_sold', self.get_language()),
                code='only_active_can_be_sold',
            )
        stock = actives[0]
        stock.status = ProductStock.Status.SOLD
        stock.save()
        _clear_public_market_cache()
        return StandardResponse.success(
            data=ProductSerializer(product).data,
            message=t('marketplace.product_sold', self.get_language())
        )

    def _set_active(self, is_active, message_key):
        """Product-level switch. Switching off pauses the live batches;
        switching back on restores exactly those (see set_product_active).
        A product whose batches are all paused, none by the switch-off, stays
        inactive — the FPO activates a batch, which switches it on."""
        product = self.get_object()
        if is_active and not product.is_active and activate_would_leave_all_inactive(product):
            raise BusinessLogicError(
                message=t('marketplace.product_all_batches_inactive', self.get_language()),
                code='product_all_batches_inactive',
            )
        if set_product_active(product, is_active, restore_batches=True):
            _clear_public_market_cache()
            # Re-read: the batch statuses changed under the prefetched `stocks`.
            product = self.get_object()
        return StandardResponse.success(
            data=ProductSerializer(product, context=self.get_serializer_context()).data,
            message=t(message_key, self.get_language()),
        )

    @extend_schema(tags=['Marketplace - Products'], summary='Hide the product from buyers and pause its live batches')
    @action(detail=True, methods=['post'])
    def deactivate(self, request, pk=None):
        return self._set_active(False, 'marketplace.product_deactivated')

    @extend_schema(tags=['Marketplace - Products'], summary='Show the product to buyers again and restore the batches it paused')
    @action(detail=True, methods=['post'])
    def activate(self, request, pk=None):
        return self._set_active(True, 'marketplace.product_activated')

    @extend_schema(tags=['Marketplace - Products'], summary='Download bulk-import Excel template')
    @action(detail=False, methods=['get'], url_path='bulk-template')
    def bulk_template(self, request):
        """
        GET /api/marketplace/products/bulk-template/

        Downloads an .xlsx with three sheets:
          - "Products"        — the sheet the FPO fills in and re-uploads
          - "Instructions"     — column meanings and rules
          - "Commodity Codes"  — every valid commodity code, for reference
        """
        wb = openpyxl.Workbook()

        # ── Sheet 1: Products (what the FPO fills in) ──
        ws = wb.active
        ws.title = 'Products'
        headers = [
            'name_en', 'name_ml', 'commodity_code', 'description_en', 'description_ml',
            'quantity', 'unit', 'price_per_unit', 'quality_certification',
            'available_from', 'available_until', 'is_public',
        ]
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        center = Alignment(horizontal='center', vertical='center')
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20
        ws.row_dimensions[1].height = 22

        # ── Sheet 2: Instructions ──
        ws2 = wb.create_sheet(title='Instructions')
        instructions = [
            ('HOW TO FILL THIS FILE', True),
            ('', False),
            ('1. Fill one row per batch, starting from row 2 of the "Products" sheet.', False),
            ('2. Do NOT change the column headers in row 1.', False),
            ('3. Required: name_en, commodity_code, description_en, quantity, unit, price_per_unit, available_from.', False),
            ('4. All other columns are optional and can be left blank.', False),
            ('5. Save as .xlsx before uploading.', False),
            ('', False),
            ('COLUMN MEANINGS', True),
            ('', False),
            ('name_en                -> Product name in English (required)', False),
            ('name_ml                -> Product name in Malayalam', False),
            ('commodity_code         -> Must match a code from the "Commodity Codes" sheet (required)', False),
            ('description_en         -> Product description in English (required)', False),
            ('description_ml         -> Product description in Malayalam', False),
            ('quantity               -> A positive number (required)', False),
            ('unit                   -> One of: kg, quintal, mt, litre, piece (required)', False),
            ('price_per_unit         -> A positive number (required)', False),
            ('quality_certification  -> Free text, e.g. FSSAI, NPOP Organic', False),
            ('available_from         -> Date in YYYY-MM-DD format (required)', False),
            ('available_until        -> Date in YYYY-MM-DD format, leave blank if open-ended', False),
            ('is_public              -> Yes/No — show on the public Market Hub (default Yes if blank)', False),
            ('', False),
            ('ONE PRODUCT, SEVERAL BATCHES', True),
            ('', False),
            ('Rows with the same name_en and commodity_code are ONE product — each row becomes a separate batch', False),
            ('of it, with its own quantity, unit, price, certification and dates.', False),
            ('The product name and description are taken from the first of those rows.', False),
            ('If you already have a product with that name and commodity code, the rows are added to it as new batches.', False),
            ('A row that exactly repeats another batch of the same product is skipped and reported as a duplicate.', False),
            ('', False),
            ('PRODUCTS SAVED AS DRAFT', True),
            ('', False),
            ('If a new product\'s name, commodity_code and description are fine but its batch columns are not', False),
            ('(e.g. quantity missing, dates wrong), the product is still created — switched off (Inactive) and with', False),
            ('no batch — and listed as a draft. Add the batch from Manage Batches; the product switches on once it is live.', False),
            ('', False),
            ('All batches imported from this sheet go live (Active) immediately.', True),
            ('Product photos cannot be added here — add them by editing each product afterward.', True),
        ]
        for row_idx, (text, is_heading) in enumerate(instructions, start=1):
            cell = ws2.cell(row=row_idx, column=1, value=text)
            if is_heading:
                cell.font = Font(bold=True, size=12)
        ws2.column_dimensions['A'].width = 95

        # ── Sheet 3: Commodity Codes (reference only) ──
        ws3 = wb.create_sheet(title='Commodity Codes')
        ws3.cell(row=1, column=1, value='code').font = Font(bold=True)
        ws3.cell(row=1, column=2, value='name').font = Font(bold=True)
        commodities = MasterLookup.objects.filter(
            category='commodity', is_active=True
        ).order_by('display_order', 'code')
        for row_idx, commodity in enumerate(commodities, start=2):
            ws3.cell(row=row_idx, column=1, value=commodity.code)
            ws3.cell(row=row_idx, column=2, value=commodity.get_name('en'))
        ws3.column_dimensions['A'].width = 20
        ws3.column_dimensions['B'].width = 35

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="product_bulk_import_template.xlsx"'
        return response

    @extend_schema(tags=['Marketplace - Products'], summary='Bulk-create products from an uploaded Excel/CSV file')
    @action(detail=False, methods=['post'], url_path='bulk-import', parser_classes=[MultiPartParser])
    def bulk_import(self, request):
        """
        POST /api/marketplace/products/bulk-import/

        Upload a filled bulk-template .xlsx (or .csv) as multipart form
        field "file". Each valid row becomes an Active batch. Rows sharing
        name_en (case/spacing ignored) + commodity_code are one product:
        the first such row creates it (its name and description win) and
        the rest are added as batches. A product the FPO already has with
        that name + commodity gets the batches instead, so re-importing
        never creates a duplicate product. A row whose batch values repeat
        an earlier row or an existing batch of the same product is skipped
        as a duplicate. A row with valid product details but invalid batch
        details still creates its product — switched off, with no batch —
        and is reported under `drafts`. Invalid rows are skipped and
        reported back so the FPO can fix just those and re-upload. Mirrors
        the pattern used by the team bulk-invite-file endpoint
        (apps/fpo/api/team.py).
        """
        fpo = get_member_fpo(request.user)

        file = request.FILES.get('file')
        if not file:
            raise BusinessLogicError(
                message='No file uploaded. Send file as multipart form field "file".',
                code='no_file_uploaded',
            )

        filename = file.name.lower()
        rows = []
        try:
            if filename.endswith('.xlsx'):
                wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
                ws = wb['Products'] if 'Products' in wb.sheetnames else wb.active
                headers = [str(c.value).strip().lower() if c.value else '' for c in next(ws.iter_rows())]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or all(v is None or str(v).strip() == '' for v in row):
                        continue  # skip fully empty rows
                    rows.append(dict(zip(headers, row)))
            elif filename.endswith('.csv'):
                content = file.read().decode('utf-8-sig')
                reader = csv.DictReader(io.StringIO(content))
                rows = [r for r in reader if any((v or '').strip() for v in r.values())]
            else:
                raise BusinessLogicError(
                    message='Only .xlsx and .csv files are supported.',
                    code='unsupported_file_type',
                )
        except BusinessLogicError:
            raise
        except Exception as e:
            raise BusinessLogicError(
                message=f'Could not parse file: {e}',
                code='file_parse_error',
            )

        if not rows:
            raise BusinessLogicError(
                message='No rows to import — the file contains only headers '
                        'or empty rows.',
                code='empty_file',
            )

        # Tester 2026-10-12: a 500-row xlsx took ~58s on the test server —
        # past every client timeout — and the browser reported 'Upload
        # failed' while the server quietly kept importing. Cap the batch at
        # a size that completes comfortably within request timeouts and
        # tell the user the limit instead of timing out.
        MAX_BULK_IMPORT_ROWS = 100
        if len(rows) > MAX_BULK_IMPORT_ROWS:
            raise BusinessLogicError(
                message=f'Too many rows: this file has {len(rows)} data '
                        f'rows but a single import is limited to '
                        f'{MAX_BULK_IMPORT_ROWS}. Split the file and '
                        f'upload in parts.',
                code='too_many_rows',
            )

        success, failed, created_batches = [], [], []
        # One commodity lookup per distinct code instead of one per row.
        commodity_cache: dict = {}
        # Products are keyed by (normalized name_en, commodity id). Seeded
        # with the FPO's existing products (oldest wins if it already has
        # duplicates) so their rows become batches instead of new products.
        products_by_key: dict = {}
        for existing in Product.objects.filter(fpo=fpo, is_deleted=False).order_by('created_at'):
            name = existing.name.get('en', '') if isinstance(existing.name, dict) else existing.name
            products_by_key.setdefault((_normalize_product_name(name), existing.commodity_id), existing)
        # Per product key: batch signature -> duplicate reason. Loaded
        # lazily from an existing product's batches on its first row.
        seen_batches: dict = {}
        new_product_ids, updated_product_ids = set(), set()
        # Products this file created without a batch (switched off because
        # their row's batch details failed): product id -> draft entry. If a
        # later row adds a batch, the product is switched on (as is any
        # switched-off product that gets a batch) and the entry moves to
        # `failed`, since only that row's batch is still missing.
        drafts_by_product: dict = {}
        # Single transaction: one commit instead of per-row fsyncs (the
        # dominant cost on the small test-server Postgres), and an
        # unexpected crash mid-file leaves nothing half-imported. Per-row
        # validation failures are still caught and skipped as before.
        with transaction.atomic():
            for i, row in enumerate(rows, start=2):  # row 1 is the header
                try:
                    data = _parse_bulk_product(row, commodity_cache)
                    key = (_normalize_product_name(data['name_en']), data['commodity'].id)
                    product = products_by_key.get(key)

                    try:
                        batch_fields, batch_error = _parse_bulk_batch(row), None
                    except ValueError as e:
                        batch_fields, batch_error = None, str(e)
                    if batch_fields is None and product is not None:
                        # The product already exists, so this row has
                        # nothing left to create.
                        raise ValueError(batch_error)

                    if batch_fields is not None:
                        seen = seen_batches.get(key)
                        if seen is None:
                            seen = {}
                            if product is not None:
                                for stock in product.stocks.filter(is_deleted=False):
                                    seen.setdefault(
                                        _batch_signature(stock),
                                        f'Duplicate — this batch already exists ({stock.batch_number})',
                                    )
                            seen_batches[key] = seen
                        signature = _batch_signature(batch_fields)
                        if signature in seen:
                            raise ValueError(seen[signature])

                    # Per-row savepoint so a DB-level failure on one row
                    # (e.g. IntegrityError) is skipped without poisoning
                    # the outer transaction.
                    is_new_product = product is None
                    batch = None
                    with transaction.atomic():
                        if is_new_product:
                            # A bad batch still creates the product, switched
                            # off and without a batch, for the FPO to finish.
                            product = Product.objects.create(
                                fpo=fpo,
                                name={'en': data['name_en'], 'ml': data['name_ml']},
                                commodity=data['commodity'],
                                description={'en': data['description_en'], 'ml': data['description_ml']},
                                is_active=batch_fields is not None,
                            )
                        if batch_fields is not None:
                            batch = ProductStock.objects.create(
                                product=product,
                                status=ProductStock.Status.ACTIVE,
                                **batch_fields,
                            )
                            # Covers drafts from this file or an earlier
                            # upload, and products the FPO switched off.
                            switch_product_on_for_live_batch(product)
                    # Only remembered once the savepoint has committed, so a
                    # rolled-back product is never reused by a later row.
                    products_by_key[key] = product
                    (new_product_ids if is_new_product else updated_product_ids).add(product.id)
                    if batch is None:
                        drafts_by_product[product.id] = {
                            'row': i,
                            'product_id': product.id,
                            'name': product.name.get('en', ''),
                            'reason': batch_error,
                        }
                        continue
                    draft = drafts_by_product.pop(product.id, None)
                    if draft:
                        failed.append({'row': draft['row'], 'name_en': draft['name'], 'reason': draft['reason']})
                    seen[signature] = f'Duplicate of row {i} — same product and batch details'
                    created_batches.append(batch)
                    success.append({
                        'row': i,
                        'product_id': product.id,
                        'name': product.name.get('en', ''),
                        'batch_number': batch.batch_number,
                        'new_product': is_new_product,
                    })
                except Exception as e:
                    failed.append({
                        'row': i,
                        'name_en': str(row.get('name_en', '') or ''),
                        'reason': str(e),
                    })

        if created_batches:
            for batch in created_batches:
                run_matching(batch)
            _clear_public_market_cache()

        # A product created by this file and given more batches by later
        # rows counts as new only; one still without a batch is a draft.
        updated_product_ids -= new_product_ids
        products_created = len(new_product_ids - drafts_by_product.keys())
        drafts = sorted(drafts_by_product.values(), key=lambda d: d['row'])
        failed.sort(key=lambda e: e['row'])
        return StandardResponse.success(
            data={
                'success': len(success),
                'drafted': len(drafts),
                'failed': len(failed),
                'products_created': products_created,
                'products_updated': len(updated_product_ids),
                'results': success,
                'drafts': drafts,
                'errors': failed,
            },
            message=f'{len(success)} batch(es) imported into {products_created} new and '
                    f'{len(updated_product_ids)} existing product(s), {len(drafts)} product(s) '
                    f'saved as draft without a batch, {len(failed)} failed.',
        )
