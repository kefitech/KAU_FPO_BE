"""
Generate the KAU AI-response DOCX — matrix mapping every concern from
KAU's Documents/DPR-RESPONSE/AI.docx review against what we shipped.

Uses the same KAU green styling as scripts/generate_dpr_user_stories_docx.py
so it reads consistently with the rest of the KAU-facing artefacts.

Output: Documents/DPR-RESPONSE/AI-response.docx

Run:
    source venv/bin/activate && python scripts/generate_ai_response_docx.py

Author: Athul Gopan (Kefi Tech Solutions)
"""
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

KAU_GREEN = '2e7d32'
KAU_PALE_GREEN = 'f1f8e9'
KAU_GREEN_RGB = RGBColor(0x2E, 0x7D, 0x32)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x55, 0x55, 0x55)


def _shade(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill_hex)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def _run(paragraph, text, *, bold=False, italic=False, size=10, color=None):
    r = paragraph.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def _section_header(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    _run(p, text, bold=True, size=14, color=KAU_GREEN_RGB)


def _para(doc, text, *, italic=False, size=10):
    p = doc.add_paragraph()
    _run(p, text, italic=italic, size=size)
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    _run(p, text, size=10)
    return p


def _response_table(doc, rows):
    """3-column table: KAU concern / What we shipped / Evidence.
    Green header, white bold text, pale-green alt rows."""
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.autofit = False
    widths = [5.5, 7.5, 3.5]
    for i, w in enumerate(widths):
        for r in table.rows:
            r.cells[i].width = Cm(w)

    for i, h in enumerate(['KAU concern (AI.docx)', 'What we shipped', 'Evidence']):
        cell = table.rows[0].cells[i]
        cell.text = ''
        _run(cell.paragraphs[0], h, bold=True, size=10, color=WHITE)
        _shade(cell, KAU_GREEN)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for r_idx, (concern, action, evidence) in enumerate(rows, start=1):
        fill = KAU_PALE_GREEN if r_idx % 2 == 1 else None
        for c_idx, val in enumerate([concern, action, evidence]):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ''
            first = True
            for line in val.split('\n'):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                _run(p, line, size=9)
                first = False
            if fill:
                _shade(cell, fill)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    doc.add_paragraph('')


def _title_block(doc):
    p = doc.add_paragraph()
    _run(p, 'KAU–FPO Linkage Platform', bold=True, size=18, color=KAU_GREEN_RGB)
    p = doc.add_paragraph()
    _run(p, 'AI Narrative Review — Response Matrix',
         bold=True, size=14, color=KAU_GREEN_RGB)
    p = doc.add_paragraph()
    _run(p, 'Kerala Agricultural University  ·  Communication Centre', size=10)
    p = doc.add_paragraph()
    _run(p, 'Prepared by: Athul Gopan · Kefi Tech Solutions  ·  19 September 2026',
         italic=True, size=10)
    doc.add_paragraph('')


# ─────────────────────────────────────────────────────────────────────────────
# Response rows — one per KAU concern from Documents/DPR-RESPONSE/AI.docx
# ─────────────────────────────────────────────────────────────────────────────

CROSS_CUTTING_ROWS = [
    (
        'Kill unresolved placeholders — [X MT per day], [Rs. X Lakhs], '
        '[CIN Number], [Name of the CEO] etc. must never appear.',
        'Placeholder scrubber runs after every LLM call. Regex catches every '
        'shape KAU flagged (bracket-X-amount, [Name of …], [CIN Number], '
        '[Location/Taluk], [projected …], [TODO …], [Placeholder …], and '
        '10+ more) and replaces each with "Not available". Preserves our own '
        '[KB #N] citations via negative lookahead. Chapter is flagged '
        'needs_review=true and the raw hits are stored on '
        'DPRAIContent.placeholder_hits so the FPO can go back and fill them in.',
        'apps/fpo/services/dpr/narrative.py — scrub_placeholders(); '
        'model field DPRAIContent.placeholder_hits, migration 0099. '
        'Regenerated sample DPR shows 0 hits across 11 chapters '
        '(before-grounding.pdf vs after-all-phases.pdf).',
    ),
    (
        'AI must not invent numbers — every fact must trace to (a) FPO input, '
        '(b) calc-engine output, (c) verified knowledge base, or '
        '(d) AI interpretation of the above.',
        'FACTS block injected into every prompt. Built directly from '
        'CalculationResult (project cost, MoF breakdown, DE ratio, Y1 revenue/'
        'opex/EBITDA/PAT, IRR, NPV, DSCR avg + min, payback, break-even). '
        'HARD RULE 7 forbids the LLM from inventing / estimating / '
        'approximating any number; HARD RULE 8 forbids bracketed placeholders; '
        'HARD RULE 9 forbids fabricating FPO facts (certifications, buyers, '
        'turnover, board composition).',
        'apps/fpo/services/dpr/narrative.py — format_calc_facts_for_prompt(), '
        '_HARD_RULES constant. 26/26 unit tests including determinism check.',
    ),
    (
        'AI generates AFTER calculations complete. Every AI narrative must '
        'read the final validated calc-engine values, not re-estimate them.',
        'generate_chapter() and generate_all_narratives() call compute(project) '
        'FIRST and refuse to run if the calc engine fails (NarrativeError). '
        'generate_all_narratives() computes ONCE and passes the same facts '
        'string into every chapter so all 11 chapters see the identical values.',
        'apps/fpo/services/dpr/narrative.py — order-enforcement path. '
        'Unit tests: test_refuses_when_compute_raises, '
        'test_computes_once_and_reuses_facts_across_chapters.',
    ),
    (
        'Cross-chapter reconciliation — the same figure must never appear '
        'differently in different chapters.',
        'Post-generation consistency check extracts numeric mentions around '
        'metric keywords (project cost, means of finance, debt:equity, IRR, '
        'NPV, DSCR avg + min, payback, Y1 PAT) from each chapter and compares '
        'against the CalculationResult with a per-metric tolerance. Drift '
        '(< 2%) vs mismatch (> 2%) classified; multiple valid candidate values '
        'accepted per metric (e.g. quoting min DSCR OR avg DSCR is fine). '
        'Warnings persisted to DPRAIContent.consistency_warnings.',
        'apps/fpo/services/dpr/consistency_check.py. Migration 0101. '
        '6 unit tests including drift/mismatch classification. '
        'Regenerated Wayanad-style sample: 0 warnings across 11 chapters.',
    ),
    (
        'Distinguish user-provided vs system-default vs AI-inferred values. '
        'The platform should identify which numbers are KAU-configured '
        'assumptions vs project-specific inputs.',
        'Every DPRConfig-configured rate (discount rate 12%, corporate tax '
        '25.17%, inflation 5%, loan-interest fallback, 3 depreciation rates, '
        'cost-variance tolerance) is tagged [system_default] in the FACTS '
        'block AND enumerated in a new "Key Assumptions Used" mini-table on '
        'the PDF, rendered just before Limitations. Every occurrence in the '
        'DPR body of a system-default rate carries an orange asterisk (e.g. '
        '"12.00%*") pointing back to the Assumptions table. HARD RULE 10 '
        'requires the LLM to indicate provenance in prose whenever it quotes '
        'a [system_default] value.',
        'apps/fpo/services/dpr/provenance.py, apps/fpo/services/dpr/pdf.py '
        '(_key_assumptions_rows), apps/fpo/templates/dpr/report.html '
        '(sec-assumptions + sup.prov CSS). Live-tested: Financial Analysis '
        'chapter cites "at the platform\'s default 12% discount rate", '
        '"platform-default corporate tax rate of 25.17%", '
        '"system-default general inflation rate of 5%".',
    ),
    (
        'Editable AI output — user must be able to review + edit AI text '
        'before final DPR generation.',
        'Every chapter card on /fpo/dpr/<uuid>/ai-content has "Edit in place" '
        'mode with textarea. Explicit Save button + autosave-on-blur (fires '
        'when the text differs from the last-saved user_edited value). '
        'Backend PATCH endpoint at /api/fpo/dpr/<uuid>/ai-content/<chapter>/ '
        'writes to user_edited. All three storage layers (DPRAIContent slots, '
        'AIUsageLog for cost, AuditLog for full trail) capture the edit.',
        'src/app/fpo/(portal)/dpr/[uuid]/ai-content/page.tsx — '
        'editMutation + onBlur handler. Existing accept/keep/merge flows '
        'preserved. FE contract at src/lib/api/dpr-ai-content.ts.',
    ),
    (
        'Debt:Equity ratio should be a simplified ratio (e.g. 5.00:1), not '
        'raw rupee amounts.',
        'Now renders as "x.xx : 1" (banking convention) everywhere — FACTS '
        'block, Project-at-a-Glance row 11 on the PDF, Financial Analysis '
        'narrative. Divide-by-zero safe: renders "Not available" when '
        'promoter contribution is zero.',
        'apps/fpo/services/dpr/pdf.py — _debt_equity_ratio_display(). '
        'apps/fpo/templates/dpr/report.html row 11 uses '
        '{{ debt_equity_ratio_display }}. Live sample shows "2.00 : 1" '
        'instead of "₹75L : ₹37.5L".',
    ),
]


SECTION_ROWS = [
    (
        'Executive Summary — Revision Required. Currently contains unresolved '
        'placeholders even though values are available in the DPR tables.',
        'Every number in Exec Summary now pulled from CalculationResult via '
        'the FACTS block. Zero placeholder tokens across 11 chapters after '
        'regen; scrubber verified. Numbers reconcile with Project-at-a-Glance '
        'because both read from the SAME CalculationResult object.',
        'after-all-phases.pdf → Executive Summary section. '
        'Regeneration cost ₹2.61 for the full 11-chapter DPR.',
    ),
    (
        'Project Background — needs to be strengthened; avoid unsupported '
        'claims like "extensive grower-member network", "urban retail consumer '
        'markets" unless backed by questionnaire data or knowledge base.',
        'HARD RULE 9 forbids the LLM from claiming certifications, buyers, '
        'awards, land, staff, or turnover that are not in the FACTS block or '
        'the knowledge base. Chapter brief updated to instruct the LLM to '
        'ground background statements in the FPO\'s actual questionnaire '
        'inputs; unsupported statements should be replaced with "Not '
        'available" in prose.',
        'apps/fpo/services/dpr/narrative.py — HARD RULE 9 in _HARD_RULES.',
    ),
    (
        'Promoter Profile — should use actual FPO/enterprise name, CEO, '
        'CIN, share capital, Board details, membership. Also add new '
        'questions for Name of CEO, Area/acreage, Project Steering Committee.',
        'Added 8 new fields on DPRProject: ceo_name, ceo_qualification, '
        'ceo_experience_years, total_area_acreage, women_shareholding_pct, '
        'landholding_summary, board_meeting_frequency, psc_members (list of '
        '{name, role, affiliation}). Wired into the DPR wizard §2.2 '
        'Identification section under a new "8. Promoter Profile detail" '
        'card. Facts block pulls all 8 plus the existing FPO.total_members / '
        'female_members / total_directors / women_directors so the narrative '
        'can quote real numbers. Chapter brief rewritten to cite these '
        'fields verbatim from the FACTS block.',
        'Migration 0100. apps/database/models/dpr/project.py + '
        'apps/fpo/api/dpr/serializers.py (FE-facing) + '
        'apps/accounts/api/admin/dpr/project_detail.py (admin-facing). '
        'FE component: src/app/fpo/(portal)/dpr/[uuid]/sections/'
        '_components/identification-section.tsx §8 card.',
    ),
    (
        'Project at a Glance — retained but must reflect final validated '
        'calc-engine outputs, no independently reconstructed figures. '
        'Debt-equity as conventional ratio (5.00:1), NOT rupee amounts.',
        'Row 11 now shows "2.00 : 1" via debt_equity_ratio_display context '
        'key. Every other row reads directly from r.<field> on '
        'CalculationResult — grep-audited in the template lint test. NPV row '
        '15 has an orange * pointing to the Key Assumptions table.',
        'apps/fpo/templates/dpr/report.html Project at a Glance section. '
        'Test: TemplateSingleSourceTests.'
        'test_template_has_no_arithmetic_filters — 0 arithmetic filters '
        'present; test would fail if anyone re-introduces inline calc.',
    ),
    (
        'Depreciation Schedule — page orientation is landscape but the table '
        'is not fully visible.',
        'New .dense-table CSS variant (6.75pt font, fixed layout, tight '
        'padding, word-wrap). Applied to Depreciation Schedule (14 columns), '
        'P&L, Cash Flow, and Balance Sheet — the four widest 11-14 column '
        'tables. Column headers shortened: "Y1 depreciation" → "Y1", '
        '"Net block Y10" → "Net Y10". Fits inside landscape A4 end-to-end.',
        'after-all-phases.pdf → §5 Depreciation Schedule.',
    ),
    (
        'Donut Pie visibility can be improved.',
        'Chart rebuilt with two-panel layout (7.5×4.5in canvas): donut on '
        'left, dedicated legend on right showing each slice as "<label> — '
        '₹ <indian_amount> (<pct>%)". No more wedge-side label collisions. '
        'In-wedge % threshold raised to ≥ 4% for cleaner rendering. Bigger '
        'title (13pt) + bigger percentage labels (9pt bold).',
        'apps/fpo/services/dpr/chart_helpers.py — cost_breakdown_pie().',
    ),
    (
        'Financial ratios (IRR, NPV, DSCR, Payback, Break-even, BCR) must be '
        'calculated from the final validated financial model and reconcile '
        'across chapters. Assumption disclosure required.',
        'All ratios computed once in calculation.py; FACTS block surfaces '
        'them; the LLM reads them and cites them by name. Consistency check '
        'catches any drift in prose citations. Assumption disclosure: '
        '"Key Assumptions Used" mini-table on the PDF lists every '
        'system-default rate the calc engine used with its source. HARD RULE '
        '10 additionally requires prose disclosure whenever the narrative '
        'quotes a system-default value.',
        'apps/fpo/services/dpr/calculation.py, provenance.py. '
        'after-all-phases.pdf → Key Assumptions Used section.',
    ),
    (
        'Risk Assessment — must be project-specific from questionnaire; '
        'avoid generic AI-generated statements. Use 3x3 (or configured NxN) '
        'probability × impact matrix. Overall risk derived by policy.',
        'Already wired end-to-end in Phase 4 (RCD B.9): DPRRiskMatrix + '
        'DPRRiskMatrixCell admin-configurable at /admin/dpr-risk-matrix, '
        'per-category worst-case class + overall rule applied by '
        'calc_engine.build_risk_assessment(). AI risk narrative reads the '
        'assessed risks from the FACTS block; HARD RULE 9 forbids inventing '
        'risks that were not entered. Consistency check does not run over '
        'text-only risk items — that\'s a Phase 4 policy path.',
        'apps/fpo/services/dpr/calculation.py — build_risk_assessment(). '
        'Admin UI: /admin/dpr-risk-matrix.',
    ),
    (
        'Limitations & Guidelines — should be project-specific and generated '
        'from the actual information gaps + estimation assumptions in the '
        'specific DPR.',
        'Two paths now feed this: (1) HARD RULE 10 forces the LLM to name '
        'every [system_default] value in prose and label them as platform '
        'assumptions rather than project facts. (2) DPRAIContent.needs_review '
        '+ DPRAIContent.consistency_warnings surface any auto-scrubbed '
        'placeholders and any numeric drift; both are shown to the FPO on '
        'the AI Content page (rose banner) and to KAU admin on the DPR '
        'oversight page (new "AI Content Health" card).',
        'Frontend surfaces: '
        'src/app/fpo/(portal)/dpr/[uuid]/ai-content/page.tsx (FPO), '
        'src/app/admin/dpr/projects/[uuid]/_components/'
        'ai-content-health.tsx (admin).',
    ),
]


TEST_COVERAGE_NOTE = (
    'End-to-end regression: 37 hermetic unit tests under '
    'apps/fpo/tests_narrative_grounding.py cover placeholder scrubber '
    '(7 tests), FACTS formatter (6), Promoter Profile fields (7), '
    'build_prompt (4), order-enforcement (2), provenance (3), consistency '
    'check (6), and template single-source lint (2). Every test runs '
    'against Python objects only — no LLM call, no DB — so the suite '
    'stays fast and hermetic. Live-Gemini validation on the Coconut Oil '
    'sample project (uuid 7a44dcf8) produces 0 placeholder hits, 0 '
    'consistency warnings, ₹2.61 total cost for the 11-chapter '
    'regeneration.'
)


def build_document(output_path: str):
    doc = Document()
    _title_block(doc)

    _para(doc,
        'This document is the response matrix to KAU\'s AI narrative review '
        'shared as Documents/DPR-RESPONSE/AI.docx (18 September 2026). Every '
        'concern raised in that review has been mapped to a specific '
        'implementation change on the platform, verified by unit tests and '
        'live-Gemini regeneration of a sample DPR.')
    _para(doc,
        'Reference DPR side by side: '
        'Documents/DPR-RESPONSE/before-grounding.pdf (baseline before this '
        'work) vs Documents/DPR-RESPONSE/after-all-phases.pdf (final).',
        italic=True)

    _section_header(doc, 'Cross-cutting concerns')
    _para(doc,
        'These themes apply to every DPR section KAU reviewed. Fixing them '
        'once at the pipeline level cascades through all 11 AI-generated '
        'chapters.')
    _response_table(doc, CROSS_CUTTING_ROWS)

    _section_header(doc, 'Section-specific concerns')
    _para(doc,
        'Where KAU flagged a specific chapter or section for revision, the '
        'per-section fix is captured below.')
    _response_table(doc, SECTION_ROWS)

    _section_header(doc, 'Testing + validation')
    _para(doc, TEST_COVERAGE_NOTE)

    _section_header(doc, 'What was intentionally deferred')
    for line in [
        'Prompt text moved to DB (DPRPromptTemplate model). KAU did not '
        'ask for admin-editable prompts; prompts stay in Python for this '
        'pass. Deferred to a later sprint — a `DPRPromptTemplate` model '
        'with admin CRUD can be added without touching the grounding '
        'pipeline.',
        'Consistency check extension to non-numeric content (buyer names, '
        'certifications, dates). Out of scope for the numeric-focused '
        'audit; HARD RULE 9 prevents fabrication of that class of fact '
        'in prose.',
    ]:
        _bullet(doc, line)

    doc.add_paragraph('')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p,
        'Prepared by Athul Gopan · Kefi Tech Solutions for KAU-FPO Linkage '
        'Programme. Ready for KAU sign-off before UAT resumes.',
        italic=True, size=9, color=MUTED)

    doc.save(output_path)
    print(f'Wrote {output_path}')


if __name__ == '__main__':
    build_document('Documents/DPR-RESPONSE/AI-response.docx')
