"""
One-shot: convert Documents/DPR_Module_For_KAU_Review.md → .docx

Uses python-docx (already in requirements). Renders headings, paragraphs,
tables and code blocks with sensible styling — nothing fancy, just clean
enough to hand to KAU leadership. Emoji/box-drawing characters in the
status-flow ASCII block render as a monospaced code paragraph.

Run:
    source venv/bin/activate
    python scripts/generate_dpr_review_docx.py
Output:
    Documents/DPR_Module_For_KAU_Review.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import sys

if len(sys.argv) > 1:
    SRC = Path(sys.argv[1])
    OUT = SRC.with_suffix(".docx")
else:
    SRC = Path("Documents/DPR_Module_For_KAU_Review.md")
    OUT = Path("Documents/DPR_Module_For_KAU_Review.docx")

# KAU green palette — matches scripts/generate_dpr_user_stories_docx.py
KAU_GREEN_RGB = RGBColor(0x2E, 0x7D, 0x32)
KAU_GREEN_HEX = "2e7d32"
KAU_PALE_GREEN_HEX = "f1f8e9"
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)


def _shade(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _style_heading(paragraph, level):
    for run in paragraph.runs:
        run.font.color.rgb = KAU_GREEN_RGB
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(20)
        elif level == 2:
            run.font.size = Pt(15)
        elif level == 3:
            run.font.size = Pt(13)
        else:
            run.font.size = Pt(11)


def _add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(11)
    return p


def _flush_table(doc, rows):
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell = table.rows[i].cells[j]
            cell.text = row[j] if j < len(row) else ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if i == 0:
                _shade(cell, KAU_GREEN_HEX)
            elif i % 2 == 1:
                _shade(cell, KAU_PALE_GREEN_HEX)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = WHITE_RGB


def _flush_code(doc, lines):
    if not lines:
        return
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.3)


INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
INLINE_CODE = re.compile(r"`([^`]+)`")


def _write_inline(paragraph, text):
    """Split a line into runs so **bold**, *italic*, `code` render properly."""
    tokens = []
    i = 0
    while i < len(text):
        # Find the earliest of bold/italic/code
        matches = []
        for pat, kind in ((INLINE_BOLD, "bold"), (INLINE_ITALIC, "italic"), (INLINE_CODE, "code")):
            m = pat.search(text, i)
            if m:
                matches.append((m.start(), m.end(), kind, m.group(1)))
        if not matches:
            tokens.append(("plain", text[i:]))
            break
        matches.sort()
        start, end, kind, inner = matches[0]
        if start > i:
            tokens.append(("plain", text[i:start]))
        tokens.append((kind, inner))
        i = end
    for kind, val in tokens:
        run = paragraph.add_run(val)
        run.font.size = Pt(11)
        if kind == "bold":
            run.bold = True
        elif kind == "italic":
            run.italic = True
        elif kind == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(10)


def main():
    text = SRC.read_text(encoding="utf-8")
    doc = Document()

    # Sensible defaults
    for style_name in ("Normal",):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

    lines = text.splitlines()

    buf_table = []
    buf_code = []
    in_code = False

    def flush_all():
        nonlocal buf_table, buf_code
        _flush_table(doc, buf_table)
        buf_table = []
        _flush_code(doc, buf_code)
        buf_code = []

    for raw in lines:
        line = raw.rstrip()

        # Code fences
        if line.startswith("```"):
            if in_code:
                _flush_code(doc, buf_code)
                buf_code = []
                in_code = False
            else:
                _flush_table(doc, buf_table)
                buf_table = []
                in_code = True
            continue
        if in_code:
            buf_code.append(raw)
            continue

        # Table rows
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            # Skip separator row like |---|---|
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            buf_table.append(cells)
            continue
        else:
            _flush_table(doc, buf_table)
            buf_table = []

        # Headings
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.add_run(line[2:].strip())
            _style_heading(p, 1)
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.add_run(line[3:].strip())
            _style_heading(p, 2)
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.add_run(line[4:].strip())
            _style_heading(p, 3)
            continue

        # Horizontal rule
        if stripped in ("---", "***"):
            doc.add_paragraph()
            continue

        # Bullet lists
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _write_inline(p, stripped[2:])
            continue
        m = re.match(r"^\d+\.\s+(.+)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _write_inline(p, m.group(1))
            continue

        # Italic-only line (final signature)
        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped.strip("*"))
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            continue

        # Empty
        if not stripped:
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        _write_inline(p, stripped)

    flush_all()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
