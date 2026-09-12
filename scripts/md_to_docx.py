"""
Convert markdown files to Word DOCX for KAU-shareable deliverables.

Handles: h1 / h2 / h3 headings, bold, italic, code, bullet lists, tables,
horizontal rules, blockquotes, plain paragraphs. Enough for the KAU reply
docs (`DPR_Config_Parameters_v1.md`, `DPR_Level2_Rule_Change_Log_v1.md`).

Run:
    exec(open('scripts/md_to_docx.py').read())
    md_to_docx('Documents/KAU_Reply_2026-09-09/DPR_Config_Parameters_v1.md')
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


_INLINE_RE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')


def _add_runs(paragraph, text):
    """Split text on inline markdown (bold / italic / code) and add runs."""
    for chunk in _INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith('**') and chunk.endswith('**'):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith('*') and chunk.endswith('*'):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        elif chunk.startswith('`') and chunk.endswith('`'):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(0x88, 0x44, 0x00)
        else:
            paragraph.add_run(chunk)


def md_to_docx(md_path):
    """Convert one markdown file to `<same-name>.docx` in the same folder."""
    md_path = Path(md_path)
    out_path = md_path.with_suffix('.docx')

    lines = md_path.read_text(encoding='utf-8').splitlines()
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    in_table = False
    table_rows: list[list[str]] = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            in_table = False
            return
        headers = table_rows[0]
        data_rows = table_rows[2:]  # skip separator row (---|---)
        table = doc.add_table(rows=1 + len(data_rows), cols=len(headers))
        table.style = 'Light Grid Accent 1'
        # header
        for j, h in enumerate(headers):
            cell = table.rows[0].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(h.strip())
            run.bold = True
        # body
        for i, row in enumerate(data_rows):
            for j, val in enumerate(row):
                if j < len(headers):
                    cell = table.rows[1 + i].cells[j]
                    cell.text = ''
                    _add_runs(cell.paragraphs[0], val.strip())
        doc.add_paragraph()  # trailing space
        table_rows = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.rstrip()

        # Frontmatter (--- at start of file with metadata) — skip block
        if i == 0 and stripped == '---':
            i += 1
            while i < len(lines) and lines[i].rstrip() != '---':
                i += 1
            i += 1
            continue

        # Table detection: line starts with |
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_rows.append([c for c in stripped.strip('|').split('|')])
            i += 1
            continue
        elif in_table:
            flush_table()

        # Horizontal rule
        if stripped in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=4)
        # Bullet list
        elif stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            _add_runs(p, stripped[2:])
        elif re.match(r'^\d+\. ', stripped):
            m = re.match(r'^\d+\. (.*)', stripped)
            p = doc.add_paragraph(style='List Number')
            _add_runs(p, m.group(1))
        # Blockquote
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            run = p.add_run(stripped[2:])
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        # Empty line
        elif not stripped:
            pass
        # Regular paragraph
        else:
            p = doc.add_paragraph()
            _add_runs(p, stripped)

        i += 1

    if in_table:
        flush_table()

    doc.save(out_path)
    return out_path


def batch(paths):
    for p in paths:
        out = md_to_docx(p)
        print(f'  {p}  →  {out}')


if __name__ == '__main__':
    import sys
    if len(sys.argv) < 2:
        print('Usage: python scripts/md_to_docx.py <file.md> [file2.md ...]')
        sys.exit(1)
    batch(sys.argv[1:])
